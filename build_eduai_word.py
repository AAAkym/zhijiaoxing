from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
import re

root = Path(r'C:\Users\33552\Desktop\project_code')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.85)
sec.right_margin = Inches(.85)
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('EduAI Pro 项目技术材料汇编').bold = True
doc.add_paragraph('基于多智能体的自适应学习系统').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

def add_md(path):
    for line in path.read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if not line: continue
        if line.startswith('#'):
            n = len(line) - len(line.lstrip('#'))
            doc.add_heading(line[n:].strip(), level=min(n, 3))
        elif line.startswith('|'):
            continue
        else:
            doc.add_paragraph(re.sub(r'^(?:[-]|\d+\.)\s*', '', line))

for name in ['EduAI-Pro-需求分析文档.md', 'EduAI-Pro-系统架构设计文档.md', 'EduAI-Pro-数据库设计文档.md']:
    add_md(root / 'docs' / name)
    doc.add_page_break()

doc.add_heading('多智能体协同流程图', level=1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(root / 'docs/assets/eduai-pro-multi-agent-flow.png'), width=Inches(6.5))
doc.add_paragraph('图 1  EduAI Pro 多智能体协同流程').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()
doc.add_heading('前端页面与阶段性验证截图', level=1)
doc.add_paragraph('截图来自前端开发服务器。采集时后端接口未启动，受保护页面显示登录拦截状态，不能作为联调成功证明。')
for fn, cap in [('01-首页.png','图 2  首页'), ('02-登录页.png','图 3  登录页'), ('03-个人中心.png','图 4  个人中心（登录拦截状态）'), ('04-学习空间.png','图 5  学习空间（登录拦截状态）'), ('05-前后端联调结果.png','图 6  联调验证记录（后端未启动）')]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(root / 'docs/screenshots' / fn), width=Inches(6.5))
    doc.add_paragraph(cap).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.save(root / 'docs/EduAI-Pro-项目技术材料汇编.docx')
